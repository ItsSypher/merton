"""Sample workbook generator tests."""

from __future__ import annotations

import pytest

pytest.importorskip("openpyxl")

from openpyxl import load_workbook

from merton.excel.sample import write_sample_workbook


def test_workbook_contains_expected_sheets(tmp_path) -> None:
    path = write_sample_workbook(tmp_path / "sample.xlsx")
    assert path.exists()
    wb = load_workbook(path)
    sheets = set(wb.sheetnames)
    assert {"Read me", "Single firm", "Portfolio", "Backtest", "Function reference"}.issubset(
        sheets
    )


def test_function_reference_has_all_formulas(tmp_path) -> None:
    from merton.excel.functions import EXCEL_FUNCTIONS

    path = write_sample_workbook(tmp_path / "sample.xlsx")
    wb = load_workbook(path)
    ws = wb["Function reference"]
    column_a = [ws.cell(row=i, column=1).value for i in range(2, 2 + len(EXCEL_FUNCTIONS))]
    expected = [row[0] for row in EXCEL_FUNCTIONS]
    assert column_a == expected


def test_single_firm_sheet_has_formula_cells(tmp_path) -> None:
    path = write_sample_workbook(tmp_path / "sample.xlsx")
    wb = load_workbook(path)
    ws = wb["Single firm"]
    formulas = [ws.cell(row=i, column=2).value for i in range(11, 17)]
    # Every cell must start with '='.
    assert all(isinstance(f, str) and f.startswith("=MERTON") for f in formulas)
